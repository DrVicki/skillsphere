from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT_DIR = Path("/home/ubuntu/webdev-static-assets/caio-activity-templates")

BLUE = "2A63BF"
DARK_BLUE = "1E3A6E"
GOLD = "F5B942"
LIGHT_BLUE = "EEF4FF"
LIGHT_GOLD = "FFF8E7"
GRAY = "4B5563"

TEMPLATES = [
    {
        "week": "Week 1",
        "filename": "CAIO-Week-1-Mandate-and-Operating-System-Workbook",
        "title": "CAIO Mandate & 30-Day Operating System Workbook",
        "subtitle": "Define the opportunity, boundaries, and operating rhythm for your CAIO path.",
        "deliverable": "A one-page opportunity statement, task suitability matrix, CAIO charter, and four-week operating agreement.",
        "sections": [
            ("1. CAIO Opportunity Statement", "Name one organization or industry you know. Write a concise statement that identifies the operational work, measurable outcome, human boundary, and the reason this matters now.", [
                "Organization / industry:", "Recurring work or process to improve:", "Measurable business outcome:", "Human judgment, approval, or risk boundary:", "One-paragraph opportunity statement:"
            ]),
            ("2. Task Suitability Matrix", "Assess three candidate tasks. Start with observable work, a named owner, and a manageable consequence if the first pilot is wrong.", [
                "Task 1 — frequency, inputs, owner, quality check, consequence of error, recommendation:",
                "Task 2 — frequency, inputs, owner, quality check, consequence of error, recommendation:",
                "Task 3 — frequency, inputs, owner, quality check, consequence of error, recommendation:",
                "Most promising first pilot and why:",
                "Task to defer and why:"
            ]),
            ("3. One-Page CAIO Charter", "Write the smallest useful mandate for an internal or fractional CAIO engagement. Be specific about authority, allies, scope, and escalation.", [
                "90-day mandate in one sentence:", "Executive sponsor:", "Process owner:", "Technical partner:", "Risk / privacy partner:", "In scope (three items):", "Out of scope (three items):", "Weekly decision rhythm and escalation rule:"
            ]),
            ("4. Your BootCamp Operating Agreement", "Choose a primary path and protect time for evidence-building. Use the reflection to name the credibility you already bring.", [
                "Primary path: employed leader / fractional practitioner / staged transition:", "Target industry and why it fits:", "Target business or process:", "Two work sessions and one weekly review block:", "Evidence folder location and naming convention:", "Reflection: What credibility or access do you already have?"
            ]),
        ],
    },
    {
        "week": "Week 2",
        "filename": "CAIO-Week-2-Proof-and-Business-Case-Workbook",
        "title": "CAIO Proof Portfolio & Business Case Workbook",
        "subtitle": "Turn a real operational problem into dated evidence, a structured diagnosis, and a decision-ready case.",
        "deliverable": "A proof portfolio artifact, shadow audit, initial risk register, and quantified results memo.",
        "sections": [
            ("1. Proof Ladder", "Build evidence before you ask for permission. Capture what you observed, what changed, who validated it, and what you learned.", [
                "Target organization / process:", "Baseline evidence available today:", "Smallest credible proof you can create this week:", "Before / after measure:", "Named reviewer or process owner:", "What this proof does not claim:"
            ]),
            ("2. Shadow Audit", "Read the work as it actually happens. Follow one process end-to-end, including handoffs, exceptions, rework, and wait time.", [
                "Process trigger:", "Inputs and source systems:", "Steps, handoffs, and decision points:", "Delay, error, duplicate work, or customer pain observed:", "Current time / volume / cost baseline:", "Process owner’s definition of a good result:"
            ]),
            ("3. Practical Risk & Governance Check", "Define the operating boundaries before the pilot grows. Keep controls proportional to the use case and consequence of error.", [
                "Data that must not enter external tools:", "Required human review point:", "Approval required before use:", "Potential failure mode:", "Mitigation or stop condition:", "Owner of the control:"
            ]),
            ("4. Results Memo", "Write a one-page recommendation that an executive can use to decide whether to continue, revise, or stop.", [
                "Problem and why it matters:", "Evidence gathered:", "Expected value (time, quality, revenue, risk, or experience):", "Pilot recommendation:", "Investment / decision requested:", "Reflection: What proof would make a skeptical leader pay attention?"
            ]),
        ],
    },
    {
        "week": "Week 3",
        "filename": "CAIO-Week-3-Responsible-Pilot-Workbook",
        "title": "CAIO Responsible Pilot & Deployment Workbook",
        "subtitle": "Plan one human-approved deployment loop from real process mapping through handoff and governance.",
        "deliverable": "A mapped process, pilot plan, build/buy decision, governance checklist, and handoff plan.",
        "sections": [
            ("1. Process Map", "Map the current process before designing a solution. Include exceptions and the person who must own the outcome.", [
                "Process objective:", "Start event and desired end state:", "Step-by-step current workflow:", "Inputs, systems, and documents:", "Exceptions and escalation points:", "Accountable process owner:"
            ]),
            ("2. Pilot Design", "Choose a narrow first use case. State the success measure, the review point, the pilot window, and the stop condition.", [
                "Pilot hypothesis:", "Users and work volume:", "What AI assists with:", "What remains human-owned:", "Baseline and success measure:", "Pilot duration and review date:", "Stop or rollback condition:"
            ]),
            ("3. Build, Buy, or Configure", "Compare options against the job to be done—not novelty. Capture integration, security, cost, support, and reversibility.", [
                "Option A:", "Option B:", "Option C:", "Recommended option and rationale:", "Integration / data dependencies:", "Security, privacy, and procurement requirements:"
            ]),
            ("4. Handoff & Operating Controls", "Design for the work to survive you. Name the owner, operating cadence, quality checks, issue route, and scale decision.", [
                "Day-to-day owner:", "User guidance / training needed:", "Quality review method:", "Issue reporting and escalation route:", "Weekly performance review:", "Scale, revise, or stop decision rule:", "Reflection: What would make this pilot safe enough to approve?"
            ]),
        ],
    },
    {
        "week": "Week 4",
        "filename": "CAIO-Week-4-90-Day-Leadership-Brief-Workbook",
        "title": "CAIO 90-Day Leadership Brief & Capstone Workbook",
        "subtitle": "Package your proof, pilot, governance posture, and next decisions into a credible CAIO engagement briefing.",
        "deliverable": "An opportunity map, assessment-first conversation guide, credibility plan, and 90-day CAIO leadership brief.",
        "sections": [
            ("1. Empty Seat Map", "Find organizations that need accountable AI leadership. Look for operational pressure, visible work, and a gap between AI interest and execution.", [
                "Target organization:", "Signals that a CAIO-shaped problem exists:", "Business process / sponsor to approach:", "Why this organization is a credible fit for you:", "Relevant evidence you can lead with:"
            ]),
            ("2. Assessment-First Conversation", "Prepare questions that diagnose the work before recommending a tool. Use the conversation to earn the right to propose a scoped engagement.", [
                "Opening question:", "Questions about process, people, and customer impact:", "Questions about data, risk, and approvals:", "Evidence you will request:", "What a useful next meeting would decide:"
            ]),
            ("3. Credibility & Protection Plan", "Protect the quality of your work and the organization’s trust. Define claims you can support, limits you will name, and the boundaries of your engagement.", [
                "Claims supported by your evidence:", "Claims you will not make yet:", "Confidentiality / data handling commitments:", "Engagement scope and exclusions:", "How you will document decisions and outcomes:"
            ]),
            ("4. 90-Day CAIO Leadership Brief", "Build the capstone briefing. The goal is a real leadership decision, not a generic AI presentation.", [
                "Executive problem statement:", "90-day outcome and measures:", "First process / pilot recommendation:", "Governance and stop conditions:", "Stakeholders and decision rhythm:", "Proof portfolio summary:", "Decision or sponsorship requested:", "Final reflection: What would make a leader trust you with the seat?"
            ]),
        ],
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "D9E4F5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        tag = f"w:{side}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 120, start: int = 120, bottom: int = 120, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    heading = doc.add_heading(text, level=level)
    heading.style.font.color.rgb = RGBColor.from_string(DARK_BLUE if level == 1 else BLUE)
    heading.style.font.name = "Aptos Display"
    return heading


def add_field_table(doc: Document, prompts: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0].cells[0]
    set_cell_shading(header, LIGHT_BLUE)
    set_cell_border(header, BLUE)
    set_cell_margins(header)
    header.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    run = header.paragraphs[0].add_run("YOUR WORKSPACE")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    set_repeat_table_header(table.rows[0])

    for index, prompt in enumerate(prompts, start=1):
        cell = table.add_row().cells[0]
        set_cell_border(cell)
        set_cell_margins(cell, top=120, bottom=960)
        paragraph = cell.paragraphs[0]
        label = paragraph.add_run(f"{index}. {prompt}")
        label.bold = True
        label.font.size = Pt(9.5)
        label.font.color.rgb = RGBColor.from_string(GRAY)
    doc.add_paragraph()


def add_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("SkillSphere · Chief AI Officer (CAIO) 4-Week BootCamp · Activity Workbook")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def make_template(template: dict) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string("374151")

    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = label.add_run(f"SKILLSPHERE · {template['week'].upper()} ACTIVITY TEMPLATE")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    title = doc.add_heading(template["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.style.font.name = "Aptos Display"
    title.style.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(template["subtitle"])
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(GRAY)

    learner_table = doc.add_table(rows=2, cols=2)
    learner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    learner_table.style = "Table Grid"
    labels = [("Learner name", ""), ("Organization / target", ""), ("Start date", ""), ("Week of", "")]
    for cell, (label_text, _) in zip([c for row in learner_table.rows for c in row.cells], labels):
        set_cell_shading(cell, LIGHT_GOLD)
        set_cell_border(cell, GOLD)
        set_cell_margins(cell, top=90, bottom=240)
        p = cell.paragraphs[0]
        run = p.add_run(f"{label_text}: ")
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        p.add_run("________________________________")

    doc.add_paragraph()
    deliverable = doc.add_paragraph()
    deliverable.paragraph_format.space_after = Pt(10)
    intro = deliverable.add_run("WEEKLY DELIVERABLE  ")
    intro.bold = True
    intro.font.color.rgb = RGBColor.from_string(BLUE)
    deliverable.add_run(template["deliverable"])

    for section_title, instruction, prompts in template["sections"]:
        add_heading(doc, section_title, 1)
        intro = doc.add_paragraph(instruction)
        intro.paragraph_format.space_after = Pt(7)
        add_field_table(doc, prompts)

    reflection = doc.add_paragraph()
    reflection.alignment = WD_ALIGN_PARAGRAPH.CENTER
    reflection_run = reflection.add_run("Before you move on: save this completed workbook in your CAIO evidence file with today’s date.")
    reflection_run.bold = True
    reflection_run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    reflection_run.font.size = Pt(9)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"{template['filename']}.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    outputs = [make_template(template) for template in TEMPLATES]
    for output in outputs:
        print(output)
